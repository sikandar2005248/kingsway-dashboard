"""
Reads the Kingsway Finance workbook and works out every figure the phone
dashboard shows. Nothing here writes to the workbook.

The workbook is copied to a temporary file before reading, so the Kingsway
Finance app and OneDrive can keep saving the original while we read.
"""
import datetime as dt
import os
import shutil
import tempfile
import threading
from collections import OrderedDict, defaultdict

import openpyxl

ACCOUNTS = ["Cash", "Bank", "Easypaisa"]

_cache = {"mtime": None, "data": None, "error": None}
_lock = threading.Lock()


# ---------------------------------------------------------------- helpers
def _num(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        try:
            return float(str(v).replace(",", "").strip())
        except ValueError:
            return 0.0


def _date(v):
    """Return 'YYYY-MM-DD' or '' for any cell value."""
    if v is None or v == "":
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)):  # Excel serial number
        try:
            return (dt.date(1899, 12, 30) + dt.timedelta(days=float(v))).isoformat()
        except (OverflowError, ValueError):
            return ""
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return dt.datetime.strptime(s[:11].strip(), fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def _text(v):
    return "" if v is None else str(v).strip()


def _rows(wb, sheet, must_have):
    """Rows of a sheet as dicts. The header row is the first row (of the
    first 10) that contains the column named in `must_have`."""
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    all_rows = list(ws.iter_rows(values_only=True))
    hdr_i = None
    for i, r in enumerate(all_rows[:10]):
        if any(_text(c) == must_have for c in r):
            hdr_i = i
            break
    if hdr_i is None:
        return []
    header = [_text(c) for c in all_rows[hdr_i]]
    out = []
    for r in all_rows[hdr_i + 1:]:
        if all(c is None or _text(c) == "" for c in r):
            continue
        d = {}
        for k, v in zip(header, r):
            if k:
                d[k] = v
        out.append(d)
    return out


# ---------------------------------------------------------------- reading
def _read_workbook(path, semester_overrides):
    tmp_dir = tempfile.mkdtemp(prefix="kwdash_")
    tmp = os.path.join(tmp_dir, "copy.xlsx")
    try:
        shutil.copy2(path, tmp)
        wb = openpyxl.load_workbook(tmp, data_only=True, read_only=True)
        try:
            raw = {
                "bsn": _rows(wb, "BSN Fee Collection", "Amount Collected (PKR)"),
                "archive": _rows(wb, "BSN Archive Fee Collection", "Amount Collected (PKR)"),
                "catb": _rows(wb, "Category B Fees", "Fee Collected (PKR)"),
                "fines": _rows(wb, "Fine Income", "Fine Amount (PKR)"),
                "expenses": _rows(wb, "Expense Tracker", "Amount (PKR)"),
                "transfers": _rows(wb, "Account Transfers", "Amount (PKR)"),
                "settings": _rows(wb, "Settings", "Financial Medium"),
                "challans": _rows(wb, "Fee Challans", "Challan No."),
                "students": _rows(wb, "Students", "Registration No."),
            }
        finally:
            wb.close()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
    return _compute(raw, semester_overrides)


def read_users(path):
    """{username: {'hash':..., 'role':...}} from the app's hidden Users sheet."""
    tmp_dir = tempfile.mkdtemp(prefix="kwdash_")
    tmp = os.path.join(tmp_dir, "copy.xlsx")
    try:
        shutil.copy2(path, tmp)
        wb = openpyxl.load_workbook(tmp, data_only=True, read_only=True)
        try:
            rows = _rows(wb, "Users", "Username")
        finally:
            wb.close()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
    users = {}
    for r in rows:
        name = _text(r.get("Username"))
        h = _text(r.get("Password hash (not the password)"))
        if name and h:
            users[name] = {"hash": h, "role": _text(r.get("Role"))}
    return users


# ---------------------------------------------------------------- computing
def _month_add(year, month, n):
    k = (month - 1) + n
    return year + k // 12, k % 12 + 1


def _compute(raw, semester_overrides):
    # ---- income (dated)
    income = []
    for r in raw["bsn"]:
        income.append({
            "date": _date(r.get("Collection Date")),
            "source": "BSN " + (_text(r.get("Fee Type")) or "Fee"),
            "receipt": _text(r.get("Receipt No.")),
            "reg": _text(r.get("Student ID / Reg. No.")),
            "student": _text(r.get("Student Name")),
            "batch": _text(r.get("Batch")),
            "semester": _text(r.get("Semester")),
            "account": _text(r.get("Payment Medium")),
            "discount": _num(r.get("Discount (PKR)")),
            "amount": _num(r.get("Amount Collected (PKR)")),
            "remarks": _text(r.get("Description / Remarks")),
        })
    for r in raw["catb"]:
        income.append({
            "date": _date(r.get("Collection Date")),
            "source": "Category B " + (_text(r.get("Fee Type")) or "Fee"),
            "receipt": _text(r.get("Receipt No.")),
            "reg": _text(r.get("Student ID / Reg. No.")),
            "student": _text(r.get("Student Name")),
            "batch": _text(r.get("Batch")),
            "semester": "",
            "account": _text(r.get("Payment Medium")),
            "discount": 0.0,
            "amount": _num(r.get("Fee Collected (PKR)")),
            "remarks": _text(r.get("Description / Remarks")),
        })
    for r in raw["fines"]:
        income.append({
            "date": _date(r.get("Collection Date")),
            "source": "Fine " + _text(r.get("Fine Type")),
            "receipt": _text(r.get("Receipt No.")),
            "reg": _text(r.get("Student ID / Reg. No.")),
            "student": _text(r.get("Student Name")),
            "batch": _text(r.get("Batch")),
            "semester": "",
            "account": _text(r.get("Payment Medium")),
            "discount": 0.0,
            "amount": _num(r.get("Fine Amount (PKR)")),
            "remarks": _text(r.get("Description / Remarks")),
        })
    dated_income = [r for r in income if r["date"]]
    undated_income = [r for r in income if not r["date"]]

    # ---- expenses, transfers, openings
    expenses = [{
        "date": _date(r.get("Expense Date")),
        "category": _text(r.get("Category")),
        "description": _text(r.get("Description")),
        "vendor": _text(r.get("Paid To / Vendor")),
        "amount": _num(r.get("Amount (PKR)")),
        "account": _text(r.get("Payment Medium")),
        "status": _text(r.get("Payment Status")),
    } for r in raw["expenses"]]
    paid = [e for e in expenses if e["status"].lower() == "paid"]

    transfers = [{
        "date": _date(r.get("Transfer Date")),
        "from": _text(r.get("From Medium")),
        "to": _text(r.get("To Medium")),
        "amount": _num(r.get("Amount (PKR)")),
        "remarks": _text(r.get("Description / Remarks")),
    } for r in raw["transfers"]]

    openings = OrderedDict((a, 0.0) for a in ACCOUNTS)
    for r in raw["settings"]:
        m = _text(r.get("Financial Medium"))
        if m:
            openings[m] = _num(r.get("Opening Balance (PKR)"))
    accounts = [a for a in openings]

    # ---- ledger with running balance
    events = []
    for r in income:
        events.append({"date": r["date"], "account": r["account"], "ord": 1, "type": "Income",
                       "detail": r["source"] + (" · " + r["student"] if r["student"] else ""),
                       "money_in": r["amount"], "money_out": 0.0})
    for t in transfers:
        events.append({"date": t["date"], "account": t["from"], "ord": 2, "type": "Transfer out",
                       "detail": "To " + t["to"] + (" · " + t["remarks"] if t["remarks"] else ""),
                       "money_in": 0.0, "money_out": t["amount"]})
        events.append({"date": t["date"], "account": t["to"], "ord": 2, "type": "Transfer in",
                       "detail": "From " + t["from"] + (" · " + t["remarks"] if t["remarks"] else ""),
                       "money_in": t["amount"], "money_out": 0.0})
    for e in paid:
        events.append({"date": e["date"], "account": e["account"], "ord": 3, "type": "Expense",
                       "detail": " · ".join(x for x in (e["category"], e["vendor"], e["description"]) if x),
                       "money_in": 0.0, "money_out": e["amount"]})
    for ev in events:  # an account named in a record but not in Settings
        if ev["account"] and ev["account"] not in openings:
            openings[ev["account"]] = 0.0
            accounts.append(ev["account"])

    ledger = []
    for a in accounts:
        bal = openings[a]
        ledger.append({"date": "", "account": a, "type": "Opening balance", "detail": "From the Settings sheet",
                       "money_in": 0.0, "money_out": 0.0, "balance": bal})
        for ev in sorted((e for e in events if e["account"] == a), key=lambda e: (e["date"], e["ord"])):
            bal += ev["money_in"] - ev["money_out"]
            ledger.append({"date": ev["date"], "account": a, "type": ev["type"], "detail": ev["detail"],
                           "money_in": ev["money_in"], "money_out": ev["money_out"], "balance": bal})

    # ---- balances per account
    balances = []
    for a in accounts:
        ev = [e for e in events if e["account"] == a]
        inc = sum(e["money_in"] for e in ev if e["type"] == "Income")
        exp = sum(e["money_out"] for e in ev if e["type"] == "Expense")
        tin = sum(e["money_in"] for e in ev if e["type"] == "Transfer in")
        tout = sum(e["money_out"] for e in ev if e["type"] == "Transfer out")
        balances.append({"account": a, "opening": openings[a], "income": inc, "expenses": exp,
                         "transfers_in": tin, "transfers_out": tout,
                         "closing": openings[a] + inc - exp + tin - tout})
    total = {"account": "Total"}
    for k in ("opening", "income", "expenses", "transfers_in", "transfers_out", "closing"):
        total[k] = sum(b[k] for b in balances)

    # ---- opening / closing by day
    dates = sorted({e["date"] for e in events if e["date"]})
    daily_balances = []
    for a in accounts:
        bal = openings[a]
        for d in dates:
            day = [e for e in events if e["account"] == a and e["date"] == d]
            mi = sum(e["money_in"] for e in day)
            mo = sum(e["money_out"] for e in day)
            daily_balances.append({"date": d, "account": a, "opening": bal, "money_in": mi,
                                   "money_out": mo, "closing": bal + mi - mo})
            bal += mi - mo

    # ---- semester start estimates (for the undated archive)
    earliest = {}
    for s in raw["students"]:
        b = _text(s.get("Batch"))
        d = _date(s.get("Admission Date"))
        if b and d and (b not in earliest or d < earliest[b]):
            earliest[b] = d
    starts = {}
    semester_starts = []
    for b in sorted(set(list(earliest) + list(semester_overrides)), key=lambda x: (len(x), x)):
        if b in semester_overrides:
            y, m = semester_overrides[b]
            basis = "Set in config.ini"
        elif b in earliest:
            y, m = int(earliest[b][:4]), int(earliest[b][5:7])
            basis = "Earliest admission " + earliest[b]
        else:
            continue
        starts[b] = (y, m)
        semester_starts.append({"batch": b, "start_month": "%d-%02d" % (y, m), "basis": basis})

    archive = []
    for r in raw["archive"]:
        archive.append({
            "batch": _text(r.get("Batch")),
            "semester": _text(r.get("Semester")),
            "date": _date(r.get("Collection Date")),
            "amount": _num(r.get("Amount Collected (PKR)")),
        })
    # undated rows from the live sheets are treated like archive rows
    for r in undated_income:
        archive.append({"batch": r["batch"], "semester": r["semester"] or "1", "date": "", "amount": r["amount"]})

    # ---- income by period
    def week_of(d):
        x = dt.date.fromisoformat(d)
        return (x - dt.timedelta(days=x.weekday())).isoformat()

    buckets = {g: defaultdict(lambda: [0.0, 0.0]) for g in ("daily", "weekly", "monthly", "yearly")}
    for r in dated_income:
        d = r["date"]
        for g, p in (("daily", d), ("weekly", week_of(d)), ("monthly", d[:7]), ("yearly", d[:4])):
            buckets[g][p][0] += r["amount"]
    unplaced = 0.0
    for r in archive:
        if r["date"]:  # an archive row that does have a date counts as recorded
            d = r["date"]
            for g, p in (("daily", d), ("weekly", week_of(d)), ("monthly", d[:7]), ("yearly", d[:4])):
                buckets[g][p][0] += r["amount"]
            continue
        st = starts.get(r["batch"])
        if not st:
            unplaced += r["amount"]
            continue
        try:
            sem = max(1, int(float(r["semester"] or 1)))
        except ValueError:
            sem = 1
        y, m = _month_add(st[0], st[1], 6 * (sem - 1))
        buckets["monthly"]["%d-%02d" % (y, m)][1] += r["amount"]
        buckets["yearly"][str(y)][1] += r["amount"]

    # fill empty months so the timeline has no gaps
    if buckets["monthly"]:
        keys = sorted(buckets["monthly"])
        y, m = int(keys[0][:4]), int(keys[0][5:7])
        ey, em = int(keys[-1][:4]), int(keys[-1][5:7])
        while (y, m) <= (ey, em):
            buckets["monthly"]["%d-%02d" % (y, m)]
            y, m = _month_add(y, m, 1)

    periods = {g: [{"period": p, "recorded": v[0], "estimated": v[1], "total": v[0] + v[1]}
                   for p, v in sorted(buckets[g].items())] for g in buckets}
    recorded_total = sum(r["amount"] for r in dated_income) + sum(r["amount"] for r in archive if r["date"])
    estimated_total = sum(r["amount"] for r in archive if not r["date"]) - unplaced

    # ---- breakdowns
    def group(rows, key, val):
        m = defaultdict(float)
        for r in rows:
            m[r[key] or "(blank)"] += r[val]
        return sorted(({"label": k, "amount": v} for k, v in m.items()), key=lambda x: -x["amount"])

    daily = []
    for d in sorted({r["date"] for r in dated_income} | {e["date"] for e in paid if e["date"]}):
        daily.append({"date": d,
                      "income": sum(r["amount"] for r in dated_income if r["date"] == d),
                      "expenses": sum(e["amount"] for e in paid if e["date"] == d)})

    arch_by = defaultdict(lambda: [0.0, 0])
    for r in raw["archive"]:
        k = (_text(r.get("Batch")), _text(r.get("Semester")))
        arch_by[k][0] += _num(r.get("Amount Collected (PKR)"))
        arch_by[k][1] += 1

    def _intkey(s):
        try:
            return int(float(s))
        except ValueError:
            return 999

    archive_by_batch = [{"label": "Batch %s · Sem %s" % k, "amount": v[0], "receipts": v[1]}
                        for k, v in sorted(arch_by.items(), key=lambda kv: (_intkey(kv[0][0]), _intkey(kv[0][1])))]

    unverified = [{
        "challan": _text(r.get("Challan No.")),
        "student": _text(r.get("Student Name")),
        "reg": _text(r.get("Registration No.")),
        "ref": _text(r.get("Bank / Transaction Ref")),
        "amount": _num(r.get("Total (PKR)")),
    } for r in raw["challans"] if _text(r.get("Status")) == "Paid-Unverified"]

    dd = sorted(r["date"] for r in dated_income)
    return {
        "period": {"first": dd[0] if dd else "", "last": dd[-1] if dd else ""},
        "kpis": {
            "income": sum(r["amount"] for r in dated_income),
            "expenses": sum(e["amount"] for e in paid),
            "net": sum(r["amount"] for r in dated_income) - sum(e["amount"] for e in paid),
            "closing": total["closing"],
            "discounts": sum(r["discount"] for r in dated_income),
            "unverified": sum(u["amount"] for u in unverified),
            "unverified_count": len(unverified),
        },
        "balances": balances,
        "balances_total": total,
        "daily_balances": daily_balances,
        "ledger": ledger,
        "periods": periods,
        "all_time": {"recorded": recorded_total, "estimated": estimated_total,
                     "total": recorded_total + estimated_total, "unplaced": unplaced},
        "semester_starts": semester_starts,
        "daily": daily,
        "by_source": group(dated_income, "source", "amount"),
        "by_category": group(paid, "category", "amount"),
        "expenses": sorted(expenses, key=lambda e: e["date"], reverse=True),
        "income": sorted(dated_income, key=lambda r: r["date"], reverse=True),
        "unverified": unverified,
        "archive_by_batch": archive_by_batch,
        "archive_count": len(raw["archive"]),
        "archive_total": sum(_num(r.get("Amount Collected (PKR)")) for r in raw["archive"]),
    }


# ---------------------------------------------------------------- cache
def get_data(path, semester_overrides=None):
    """Data for the dashboard, re-read only when the workbook has changed.
    If a read fails (the file is mid-save), the last good data is returned
    with an error note."""
    semester_overrides = semester_overrides or {}
    with _lock:
        try:
            mtime = os.path.getmtime(path)
        except OSError as e:
            if _cache["data"] is not None:
                return _cache["data"], "Workbook not found right now (%s); showing the last good read." % e
            raise
        if _cache["data"] is not None and _cache["mtime"] == mtime:
            return _cache["data"], None
        try:
            data = _read_workbook(path, semester_overrides)
        except Exception as e:  # file locked or half-written
            if _cache["data"] is not None:
                return _cache["data"], "Couldn't read the latest workbook (%s); showing the last good read." % e
            raise
        data["file_modified"] = dt.datetime.fromtimestamp(mtime).strftime("%d %b %Y, %I:%M %p")
        _cache.update(mtime=mtime, data=data, error=None)
        return data, None
